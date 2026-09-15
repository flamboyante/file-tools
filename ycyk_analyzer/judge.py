# -*- coding: utf-8 -*-
"""
judge.py —— 判据引擎：从合并表的「判据」列读规则，按标准格式求值。

判据住在**合并解析表**里（`spec/spec_merged_*.xlsx`），不在代码里 ——
**加判据 = 在 Excel 里改一格**，不用动 py。格式由 `tools/normalize_judges.py` 归一，共 8 种：

    正常值:X[,级:异常|告警]      值等于 X 即通过（多值用 | 分隔；支持 0x/0b/FF/中文枚举）
    正常范围:a~b[,级:...]        值落在 [a,b] 内即通过
    文本:词|词[,级:...]          值含这些词即命中（用于"值本身就是中文结论"的字段，如快遥 T 段）
    结构对应:帧序号=N            该消息必须出现在帧序号=N 的包里（进 frame 自证，不逐拍判）
    总体:有变化 / 总体:按历史基线  跨拍判据（V1 暂未启用，保留语义）
    无需判据 / 待定               不判 → 「暂不判据」

口诀：**判据只描述"什么算正常"，命中不了就是问题**；严重度（异常 / 告警）由 `,级:` 给。

两种判定：
    judge(...)      —— 瞬时（逐拍）：给一个值，返回 通过 / 告警 / 异常 / 暂不判据
    summarize(...)  —— 总体（跨拍）：给"异常拍数 / 总拍数"，返回一句总体结论
                       （V1 用通用口径：全程/持续/间歇/偶发，不依赖规则）

★ 读取健壮性：openpyxl 读不了（文件被腾讯文档等工具重写过样式）时，
  自动退回**直接解析 xlsx 内部 XML** —— 样式坏不影响数据。
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from spec import default_spec_path      # 解析表定位只有一处实现（见 spec.py）

OK = "通过"
WARN = "告警"
ERR = "异常"
NA = "暂不判据"

# report.py 里的数据块名 → 合并表的 sheet 名
SOURCE_TO_SHEET = {
    "T段-快遥": "标准快遥",
    "T段-慢遥": "标准慢遥51H",
    "slot0": "通道0（slot0）",
    "slot2": "通道1（slot2）",
    "slot3": "通道2（slot3）",
    "slot5": "通道3（slot5）",
}

# 各 sheet 的「参数名称」列（通道页名称在第 4 列，T 段页在第 1 列）
NAME_COL = {
    "标准快遥": 1,
    "标准慢遥51H": 1,
    "标准慢遥52H": 1,
    "标准慢遥53H": 1,
    "通道0（slot0）": 4,
    "通道1（slot2）": 4,
    "通道2（slot3）": 4,
    "通道3（slot5）": 4,
}

# 各 sheet 的「判据」列
JUDGE_COL = {
    "标准快遥": 4,
    "标准慢遥51H": 4,
    "标准慢遥52H": 4,
    "标准慢遥53H": 4,
    "通道0（slot0）": 14,
    "通道1（slot2）": 14,
    "通道2（slot3）": 14,
    "通道3（slot5）": 14,
}

# 各 sheet 的「备注/说明」列（判据没有提到的中文含义，要从这里反查原始值）
NOTE_COL = {
    "标准快遥": 3,
    "标准慢遥51H": 3,
    "标准慢遥52H": 3,
    "标准慢遥53H": 3,
    "通道0（slot0）": 13,
    "通道1（slot2）": 13,
    "通道2（slot3）": 13,
    "通道3（slot5）": 13,
}

# 备注里的「值：含义」映射，例如「0：断开；1：建链状态」「0001b：常规业务模式」
# ★ 含义部分不能含冒号 —— 否则「1：正常：0：异常」这种连贯写法会被贪婪匹配吞成一条
RE_VALUE_MEANING = re.compile(r"([0-9A-Za-z]{1,8})\s*[:：]\s*([^;；/，,\n:：]+)")

KIND_VALUE = "value"
KIND_RANGE = "range"
KIND_TEXT = "text"
KIND_STRUCT = "struct"
KIND_SUMMARY = "summary"
KIND_NONE = "none"

POS_HINTS = ("正常", "成功", "可用", "有效", "完成", "上电", "使能")
NEG_HINTS = ("不正常", "异常", "失败", "故障", "过高", "过低", "错误", "无效",
             "未接入", "中断", "损坏", "告警", "失锁", "过功率", "不可用", "未成功")


def to_number(text: str) -> Optional[float]:
    """把各种进制写法统一转成数字；转不了返回 None。**两种风格都认**：

        十进制    31、85、0、1、1.5        （裸写，默认就是十进制）
        C 风格    0x7FFFFF、0b1010        （前缀）
        汇编风格  7FFFFFH、1010B、31D      （后缀 H=十六进制 / B=二进制 / D=十进制）

    ★ 几个踩过坑的点，都要照顾到：
      · 「0b」「1b」这种**只有标记没有数字**的写法（表里大量出现）→ 按二进制 0 / 1
        （早先会落到十六进制分支，把 "0b" 解析成 11，导致「正常」被判成异常）
      · 裸写的 `31` 一律按**十进制**（不是十六进制）—— 十进制是默认
      · 纯十六进制字符（FF）才按十六进制兜底
    """
    t = str(text).strip().lower()
    if not t:
        return None

    # ① 快路径：裸写的十进制/浮点（最常见）
    try:
        return float(t)
    except ValueError:
        pass

    # ② C 风格前缀
    if t.startswith("0x"):
        try:
            return float(int(t, 16))
        except ValueError:
            return None
    if t.startswith("0b") and len(t) > 2:
        try:
            return float(int(t[2:], 2))
        except ValueError:
            pass

    # ③ 汇编风格后缀 H / B / D
    if len(t) >= 2 and t[-1] in ("h", "b", "d"):
        body, kind = t[:-1], t[-1]
        try:
            if kind == "h":
                return float(int(body, 16))
            if kind == "b":
                # 全 0/1 → 二进制；否则当十六进制兜底（避免 "1fb" 这类误读）
                if body and set(body) <= {"0", "1"}:
                    return float(int(body, 2))
                return float(int(body, 16))
            if kind == "d" and body.isdigit():
                return float(int(body))
        except ValueError:
            pass

    # ④ 纯十六进制字符（FF）才按十六进制兜底
    if all(c in "0123456789abcdef" for c in t):
        try:
            return float(int(t, 16))
        except ValueError:
            return None
    return None


def meaning_to_raw(note: str, meaning: str) -> Optional[str]:
    """从备注的「值：含义」映射里，反查某个中文含义对应的**原始值**。

    例：备注「0：断开；1：建链状态」，含义"断开" → 返回 "0"。
    这是解决"值是中文而判据是位模式"的**可靠办法** —— 不用猜词，直接查表里的映射。
    """
    text = (meaning or "").strip()
    if not text:
        return None
    for match in RE_VALUE_MEANING.finditer(note or ""):
        candidate = match.group(2).strip()
        if candidate and (candidate in text or text in candidate):
            return match.group(1)
    return None


# ---------------------------------------------------------------------------
# 单条判据
# ---------------------------------------------------------------------------


@dataclass
class JudgeSpec:
    """一条判据（对应合并表判据列的一格）。"""

    raw: str
    kind: str
    level: str = WARN
    values: List[str] = field(default_factory=list)   # 正常值 / 文本词表
    low: Optional[float] = None
    high: Optional[float] = None
    frame_seq: Optional[int] = None
    summary_kind: str = ""
    note: str = ""                                    # 表里的备注（用于反查中文含义的原始值）

    def evaluate(self, value) -> str:
        """瞬时判定：返回 通过 / 告警 / 异常 / 暂不判据。"""
        text = "" if value is None else str(value).strip()

        if self.kind == KIND_VALUE:
            if not text:
                return NA
            got = to_number(text)
            # ★ 值不是数字时（快遥 T 段的值是**解码后的中文**，而判据那格写的是位模式），
            #   按下面三步判 —— 关键是第 ② 步「查备注映射」，**不靠猜词**：
            if got is None:
                # ① 值直接等于正常值（判据本身写成中文，如 正常值:常规业务模式）
                if any(text == v.strip() for v in self.values):
                    return OK
                # ② 从备注的「值：含义」映射反查这个中文含义的原始值，再与正常值比
                #    例：值"断开" → 备注「0：断开；1：建链状态」→ 原始值 0 ≠ 正常值 1 → 命中
                raw = meaning_to_raw(self.note, text)
                if raw is not None:
                    raw_num = to_number(raw)
                    for want in self.values:
                        want_num = to_number(want)
                        if raw_num is not None and want_num is not None and abs(raw_num - want_num) < 1e-9:
                            return OK
                    return self.level
                # ③ 实在查不到映射，才退到通用词表
                if any(w in text for w in NEG_HINTS):
                    return self.level
                if any(w in text for w in POS_HINTS):
                    return OK
                return NA
            # 数字比较（覆盖 0x / 0b / FF 与十进制混写）
            for want in self.values:
                want_num = to_number(want)
                if got is not None and want_num is not None:
                    if abs(got - want_num) < 1e-9:
                        return OK
                elif text == want.strip():
                    return OK
            return self.level

        if self.kind == KIND_RANGE:
            got = to_number(text)
            if got is None or self.low is None or self.high is None:
                return NA
            return OK if self.low <= got <= self.high else self.level

        if self.kind == KIND_TEXT:
            if not text:
                return NA
            if any(w and w in text for w in self.values):
                return self.level
            # 不含负面词：如果值本身是正面词就判通过，否则留白（如模式类枚举）
            if any(w in text for w in POS_HINTS):
                return OK
            return NA

        return NA   # 结构对应 / 总体 / 无需判据 → 不参与逐拍判

    def describe(self) -> str:
        """一句话说明这条判据（给报告用）。"""
        if self.kind == KIND_VALUE:
            return "正常值 = {}".format(" 或 ".join(self.values))
        if self.kind == KIND_RANGE:
            return "正常范围 {~{}".format(self.low, self.high)
        if self.kind == KIND_TEXT:
            return "值含「{}」即判{}".format("、".join(self.values), self.level)
        if self.kind == KIND_STRUCT:
            return "必须出现在帧序号 = {} 的包里".format(self.frame_seq)
        if self.kind == KIND_SUMMARY:
            return self.summary_kind
        return "不判"


def parse_judge(raw: str, note: str = "") -> JudgeSpec:
    """把判据列的一格文字解析成 JudgeSpec；note = 该字段在表里的备注（用来反查中文含义）。"""
    text = (raw or "").strip()

    def make(kind: str, **kw) -> JudgeSpec:
        return JudgeSpec(raw=raw or "", kind=kind, note=note, **kw)

    if not text:
        return make(KIND_NONE)

    # 先摘掉 ,级:X
    level = WARN
    match = re.search(r",\s*级[:：]\s*(\S+)", text)
    if match:
        level = match.group(1).strip()
        text = text[: match.start()].strip()

    if text.startswith("无需判据") or text.startswith("待定"):
        return make(KIND_NONE)

    if text.startswith("结构对应"):
        seq = re.search(r"帧序号\s*=\s*(\d)", text)
        return make(KIND_STRUCT, frame_seq=int(seq.group(1)) if seq else None)

    if text.startswith("总体:"):
        return make(KIND_SUMMARY, summary_kind=text)

    if text.startswith("正常范围"):
        body = text.split(":", 1)[1] if ":" in text else ""
        parts = re.split(r"[~～]", body)
        if len(parts) == 2:
            return make(KIND_RANGE, level=level,
                        low=to_number(parts[0]), high=to_number(parts[1]))
        return make(KIND_NONE)

    if text.startswith("正常值"):
        body = text.split(":", 1)[1] if ":" in text else ""
        values = [v.strip() for v in body.split("|") if v.strip()]
        return make(KIND_VALUE, level=level, values=values)

    if text.startswith("文本"):
        body = text.split(":", 1)[1] if ":" in text else ""
        values = [v.strip() for v in re.split(r"[|/、]", body) if v.strip()]
        return make(KIND_TEXT, level=level, values=values)

    return make(KIND_NONE)


# ---------------------------------------------------------------------------
# 判据表
# ---------------------------------------------------------------------------


class JudgeTable:
    """全部判据，按 (sheet 名, 字段名) 索引。"""

    def __init__(self, table: Dict[Tuple[str, str], JudgeSpec]):
        self.table = table

    def __len__(self) -> int:
        return len(self.table)

    def spec(self, source: str, field_name: str) -> Optional[JudgeSpec]:
        sheet = SOURCE_TO_SHEET.get(source, source)
        return self.table.get((sheet, field_name))

    def judge(self, source: str, field_name: str, value) -> str:
        """瞬时判定入口。没判据 / 明确无需判据 → 「暂不判据」。"""
        spec = self.spec(source, field_name)
        if spec is None:
            return NA
        return spec.evaluate(value)

    def describe(self, source: str, field_name: str) -> str:
        spec = self.spec(source, field_name)
        return spec.describe() if spec else ""

    def stats(self) -> Dict[str, int]:
        out: Dict[str, int] = {}
        for spec in self.table.values():
            key = spec.kind if spec.kind != KIND_NONE else "无需判据"
            out[key] = out.get(key, 0) + 1
        return out


def summarize(alarm_count: int, total: int) -> str:
    """总体结论（跨拍，不依赖规则）：看异常拍数占整段的比例。"""
    if total <= 0 or alarm_count <= 0:
        return ""
    ratio = alarm_count / float(total)
    if alarm_count == total:
        return "全程异常"
    if ratio >= 0.9:
        return "持续异常"
    if ratio <= 0.05:
        return "偶发"
    return "间歇异常"


# ---------------------------------------------------------------------------
# 读取（openpyxl 优先，读不了就退回 XML 抢救）
# ---------------------------------------------------------------------------

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_RELS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _rows_from_xml(path: str, sheet_name: str) -> List[List[str]]:
    """直接从 xlsx 内部 XML 读一个 sheet（绕过 openpyxl 的样式解析）。"""
    with zipfile.ZipFile(path) as z:
        shared: List[str] = []
        if "xl/sharedStrings.xml" in z.namelist():
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in root.findall(_NS + "si"):
                shared.append("".join(t.text or "" for t in si.iter(_NS + "t")))

        wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rid_map = {r.get("Id"): r.get("Target") for r in rels}
        target = None
        for sh in wb_xml.iter(_NS + "sheet"):
            if sh.get("name") == sheet_name:
                target = rid_map.get(sh.get(_RELS + "id"))
                break
        if target is None:
            return []
        # rels 里的 Target 有两种写法：相对（worksheets/sheet1.xml）
        # 或绝对（/xl/worksheets/sheet1.xml，腾讯文档这类工具会这么写）—— 两种都要兼容
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target

        root = ET.fromstring(z.read(target))
        cells: Dict[int, Dict[int, str]] = {}
        max_col = 0
        for row in root.iter(_NS + "row"):
            rno = int(row.get("r"))
            row_map: Dict[int, str] = {}
            for c in row.findall(_NS + "c"):
                m = re.match(r"([A-Z]+)(\d+)", c.get("r") or "")
                if not m:
                    continue
                col = 0
                for ch in m.group(1):
                    col = col * 26 + ord(ch) - 64
                t, v = c.get("t"), c.find(_NS + "v")
                if t == "s" and v is not None:
                    val = shared[int(v.text)]
                elif t == "inlineStr":
                    is_el = c.find(_NS + "is")
                    val = "".join(x.text or "" for x in is_el.iter(_NS + "t")) if is_el is not None else ""
                else:
                    val = v.text if v is not None and v.text is not None else ""
                row_map[col] = val
                max_col = max(max_col, col)
            cells[rno] = row_map

        return [[cells.get(r, {}).get(c, "") for c in range(1, max_col + 1)]
                for r in range(1, max(cells) + 1)] if cells else []


def _sheet_rows(path: str, sheet_name: str) -> List[List[str]]:
    """读一个 sheet 的所有行（值）；openpyxl 失败时自动退回 XML。"""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            return []
        return [list(r) for r in wb[sheet_name].iter_rows(values_only=True)]
    except Exception:
        return _rows_from_xml(path, sheet_name)


def load_judges(spec_path: Optional[str] = None) -> JudgeTable:
    """读合并表里全部判据，建索引。"""
    path = spec_path or default_spec_path()
    table: Dict[Tuple[str, str], JudgeSpec] = {}

    for sheet_name, name_col in NAME_COL.items():
        judge_col = JUDGE_COL.get(sheet_name)
        if judge_col is None:
            continue
        rows = _sheet_rows(path, sheet_name)
        note_col = NOTE_COL.get(sheet_name)
        for row in rows[1:]:
            if len(row) < max(name_col, judge_col):
                continue
            name = str(row[name_col - 1] or "").strip()
            raw = str(row[judge_col - 1] or "").strip()
            note = (str(row[note_col - 1] or "").strip()
                    if note_col and len(row) >= note_col else "")
            if not name or not raw:
                continue
            table[(sheet_name, name)] = parse_judge(raw, note)

    return JudgeTable(table)
