# -*- coding: utf-8 -*-
"""
spec.py —— 把「标三-慢遥测试表」读成 Python 对象。

这个模块只干一件事：**读懂那张 Excel 规格表**，把它变成程序能用的结构。
它不碰 CSV、不做解析、不管判据 —— 那些是后面模块的事。

表是怎么组织的（实测确认）：
    每一页 sheet 叫「通道N（slotM）」，一页 = 一个 slot 的完整帧定义（合计 1184 bit = 148 字节）。

        第 1 行    表头
        然后        以「消息ID」为界切分成若干"消息"，每条消息长这样：
                      消息ID行      ← 新消息开始（2 字节）
                      消息长度行    ← 备注里写着 length=NN字节
                      若干字段行    ← 这条消息的各个字段
        末尾        一个校验行：参数名称为空、长度bit = 1184

    ★ 切分规则：**"参数名称"里只要出现"消息ID"三个字，就是新消息的开头。**

三个实测踩过的坑（都写在代码里了）：
    1. 消息ID 有三种写法：
          消息ID（type=1000）      ← type 写在参数名称里
          消息ID                  ← 裸的，type 写在"备注"里  type=1010
          BMU消息ID(t=1050)       ← 用 t= 而不是 type=
       所以 type 必须"先看名称、再兜底看备注"。
    2. 一条消息可能对应两个 type（按制式分）：
          备注写 type=1200(sc),1201(ka) —— 所以 type 存成列表。
    3. 消息之间夹着"预留 / 保留 / 慢遥帧序号"这类**不属于消息内容**的行。
       如果把它们算进消息，消息位长就对不上备注声明的 length。
       所以这里给每行标一个 role：header / field / padding / trailer 分开统计。
"""

from __future__ import annotations

import hashlib
import os
import re
import sys
from dataclasses import dataclass, field
from typing import List, Optional

from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# 常量：列号从 1 开始，与 Excel 里看到的列号一致
# ---------------------------------------------------------------------------

COL_MODULE = 1     # 模块（OM / 业务FPGA / BMU固定填写 ...）
COL_NAME = 4       # 参数名称（"消息ID"行靠它识别）
COL_DATA_TYPE = 5  # 数据类型
COL_BITS = 12      # 长度/bit（位宽）
COL_NOTE = 13      # 备注（解码规则、枚举含义、length=NN字节 都写在这里）

# 一个 slot 包的载荷总位长 = 148 字节 × 8。页末校验行就是用这个数自检的。
EXPECTED_TOTAL_BITS = 1184

# 消息头固定 4 字节：消息ID(2B) + 消息长度(2B)
MESSAGE_HEADER_BYTES = 4

# 行的角色
ROLE_HEADER = "header"    # 消息头里的字段（消息ID / 消息长度）
ROLE_FIELD = "field"      # 正常的业务字段
ROLE_PADDING = "padding"  # 填充占位，不是业务数据（表里叫"预留"/"保留"）
ROLE_TRAILER = "trailer"  # 帧尾，不属于任何消息的业务内容（"慢遥帧序号"）

PADDING_KEYWORDS = ("预留", "保留")
TRAILER_KEYWORDS = ("帧序号",)

# 「参数名称」里出现"消息ID"就是新消息
RE_HAS_MESSAGE_ID = re.compile(r"消息ID")

# type 号：名称里或备注里都可能写，且可能是 type= 也可能是 t=。
# ★ 还要支持备注里的"多 type"写法，例如 type=1200(sc),1201(ka) —— 逗号后面那个
#   数字前面没有 type= 前缀，早先的实现只抓到 1200、漏了 1201，导致游标报"未知消息ID"。
RE_TYPE_GROUP = re.compile(
    r"(?:type|t)\s*=\s*(\d+(?:\s*\([^)]*\))?(?:\s*,\s*\d+(?:\s*\([^)]*\))?)*)",
    re.IGNORECASE,
)
RE_TYPE_NUM = re.compile(r"\d+")

# 「length=21字节」→ 抓出 21
RE_LENGTH = re.compile(r"length\s*=\s*(\d+)\s*字节", re.IGNORECASE)

# 「通道1（slot2）」→ 抓出 1 和 2
RE_SHEET_NAME = re.compile(r"通道\s*(\d+)\D+?slot\s*(\d+)", re.IGNORECASE)

# sheet 名前缀，只有通道页才解析（跳过「修改记录」）
SHEET_CHANNEL_PREFIX = "通道"


# ---------------------------------------------------------------------------
# 小工具
# ---------------------------------------------------------------------------


def cell_text(value) -> str:
    """把单元格内容变成去掉首尾空格的字符串；空值返回空串。"""
    if value is None:
        return ""
    return str(value).strip()


def cell_int(value) -> int:
    """把单元格内容变成整数；取不到数字时返回 0。"""
    if isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    match = re.search(r"\d+", cell_text(value))
    return int(match.group()) if match else 0


def _cell_at(row, col_no: int) -> str:
    """取一行里第 col_no 列的内容（列号从 1 开始），越界当空值。"""
    index = col_no - 1
    if index < 0 or index >= len(row):
        return ""
    return cell_text(row[index])


def extract_type_ids(text: str) -> List[str]:
    """从一段文字里抓出所有 type 号。

    支持三种写法：
        type=1000                → ["1000"]
        t=1050                   → ["1050"]      （BMU消息ID(t=1050)）
        type=1200(sc),1201(ka)   → ["1200", "1201"]  （一条消息两个制式）
    """
    found: List[str] = []
    for group in RE_TYPE_GROUP.findall(text or ""):
        found.extend(RE_TYPE_NUM.findall(group))
    return found


def guess_role(name: str, is_header: bool) -> str:
    """判断这一行是什么角色。"""
    if is_header:
        return ROLE_HEADER
    if any(key in name for key in PADDING_KEYWORDS):
        return ROLE_PADDING
    if any(key in name for key in TRAILER_KEYWORDS):
        return ROLE_TRAILER
    return ROLE_FIELD


SPEC_FILE_NAME = "spec_merged_20260914.xlsx"
SPEC_ENV_VAR = "YCYK_SPEC"     # 设了就优先用它；现场换表 / 新旧表对拍时用


def spec_search_dirs() -> List[str]:
    """解析表的搜索目录，按优先级排列。

    1. **程序目录** —— 打包后指 exe 所在目录。表放这儿就 **改完即生效、不用重新打包**；
       PyInstaller 单文件模式会把内容解到 `sys._MEIPASS`，也在本项里兜底
    2. **本模块目录** —— 开发态就是这里（`ycyk_analyzer/spec/`）
    """
    dirs: List[str] = []
    if getattr(sys, "frozen", False):
        dirs.append(os.path.dirname(os.path.abspath(sys.executable)))
        meipass = getattr(sys, "_MEIPASS", None)
        if meipass:
            dirs.append(meipass)
    dirs.append(os.path.dirname(os.path.abspath(__file__)))

    seen, out = set(), []                    # 去重但保持顺序
    for path in dirs:
        key = os.path.normcase(os.path.abspath(path))
        if key not in seen:
            seen.add(key)
            out.append(path)
    return out


def default_spec_path() -> str:
    """解析表路径 —— **外置优先**。

    合并解析表是本项目唯一依赖的 xlsx（slot 定义 + T 段定义 + 判据列），
    由 `tools/merge_spec_tables.py` 生成；第一页「说明」记着来源 md5 与用表须知。

    查找顺序（第一个存在的即为结果）：
      1. 环境变量 `YCYK_SPEC` 指定的文件 —— 明确指定时**不再回退**，找不到就报错，
         免得"以为换了表、其实还在用旧的"
      2. `程序目录/spec/spec_merged_*.xlsx` —— 打包分发时的**外置表**，
         改一格判据不必重新打包
      3. `本模块目录/spec/spec_merged_*.xlsx` —— 开发态，或随包打进去的兜底副本

    都没有时返回最后一个候选，让 openpyxl 抛出并带上路径，便于排查。
    """
    override = os.environ.get(SPEC_ENV_VAR)
    if override:
        return os.path.abspath(override)

    candidates = [os.path.join(d, "spec", SPEC_FILE_NAME) for d in spec_search_dirs()]
    for path in candidates:
        if os.path.isfile(path):
            return path
    return candidates[-1]


def spec_fingerprint(spec_path: Optional[str] = None) -> str:
    """解析表的「文件名 + md5 前 12 位」—— 写进报告首页，追溯这份报告用的是哪版表。"""
    path = spec_path or default_spec_path()
    try:
        with open(path, "rb") as fh:
            digest = hashlib.md5(fh.read()).hexdigest()[:12]
    except Exception:
        return "{}（读取失败）".format(os.path.basename(path))
    return "{}　md5:{}".format(os.path.basename(path), digest)


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------


@dataclass
class FieldSpec:
    """表里的一行定义。

    Attributes:
        name:   参数名称（中文字段名）
        bits:   位宽（bit）
        module: 所属模块
        note:   备注原文（判据规则的主要来源）
        row_no: 在原表里的行号，方便回查
        role:   角色，见 ROLE_* 常量
    """

    name: str
    bits: int
    module: str
    note: str
    row_no: int
    role: str = ROLE_FIELD

    @property
    def is_header(self) -> bool:
        return self.role == ROLE_HEADER


@dataclass
class MessageSpec:
    """一条消息（从"消息ID"行 到 下一条"消息ID"行之前）。"""

    type_ids: List[str] = field(default_factory=list)  # 类型号，可能 0/1/2 个
    type_raw: str = ""                                 # type 原文，便于回查
    fields: List[FieldSpec] = field(default_factory=list)
    declared_bytes: Optional[int] = None               # 备注声明的载荷字节数（不含 4 字节头）
    row_no: int = 0

    # --- 各种位长统计（分开算，才看得出规律）---

    def _bits_of(self, role: str) -> int:
        return sum(f.bits for f in self.fields if f.role == role)

    @property
    def field_bits(self) -> int:
        """业务字段位长（不含消息头、不含填充/帧尾）。"""
        return self._bits_of(ROLE_FIELD)

    @property
    def header_bits(self) -> int:
        return self._bits_of(ROLE_HEADER)

    @property
    def padding_bits(self) -> int:
        return self._bits_of(ROLE_PADDING)

    @property
    def trailer_bits(self) -> int:
        return self._bits_of(ROLE_TRAILER)

    @property
    def total_bits(self) -> int:
        """这条消息名下所有行的位长合计（含填充/帧尾，用于和页末校验对齐）。"""
        return sum(f.bits for f in self.fields)

    @property
    def declared_bits(self) -> Optional[int]:
        """按备注声明的 length 换算的位长（含 4 字节消息头）；没写 length 则 None。"""
        if self.declared_bytes is None:
            return None
        return (self.declared_bytes + MESSAGE_HEADER_BYTES) * 8

    @property
    def delta_bits(self) -> Optional[int]:
        """「业务字段+消息头」与「声明长度」的差。

        = 0   完全吻合
        != 0  表定义与声明不一致，解析会错位，必须查
        """
        if self.declared_bits is None:
            return None
        return (self.header_bits + self.field_bits) - self.declared_bits

    @property
    def payload_fields(self) -> List[FieldSpec]:
        """真正的业务字段（去掉消息头、填充、帧尾）。"""
        return [f for f in self.fields if f.role == ROLE_FIELD]

    @property
    def type_label(self) -> str:
        """给报告用的类型标签。"""
        if self.type_ids:
            return "/".join(self.type_ids)
        return "未标注"


@dataclass
class SlotSpec:
    """一页 sheet = 一个 slot 的完整帧定义。"""

    sheet_name: str
    channel_index: int
    slot_no: Optional[int]
    messages: List[MessageSpec] = field(default_factory=list)
    declared_total_bits: Optional[int] = None  # 页末校验行的值（正常 1184）

    @property
    def total_bits(self) -> int:
        """所有消息名下所有行的位长合计。"""
        return sum(m.total_bits for m in self.messages)

    @property
    def check_ok(self) -> bool:
        """自检：字段合计 = 页末校验值 = 1184。"""
        return (
            self.declared_total_bits == EXPECTED_TOTAL_BITS
            and self.total_bits == EXPECTED_TOTAL_BITS
        )


# ---------------------------------------------------------------------------
# 核心：把一页 sheet 读成 SlotSpec
# ---------------------------------------------------------------------------


def parse_sheet(rows, sheet_name: str) -> SlotSpec:
    """把一页 sheet 的所有行解析成一个 SlotSpec。"""
    name_match = RE_SHEET_NAME.search(sheet_name)
    slot = SlotSpec(
        sheet_name=sheet_name,
        channel_index=int(name_match.group(1)) if name_match else -1,
        slot_no=int(name_match.group(2)) if name_match else None,
    )

    current: Optional[MessageSpec] = None

    for row_no, row in enumerate(rows, start=1):
        if row_no == 1:  # 表头行
            continue

        name = _cell_at(row, COL_NAME)
        note = _cell_at(row, COL_NOTE)
        module = _cell_at(row, COL_MODULE)
        bits = cell_int(row[COL_BITS - 1] if len(row) >= COL_BITS else None)

        # 没有参数名称的行：空行，或页末校验行
        if not name:
            if bits == EXPECTED_TOTAL_BITS:
                slot.declared_total_bits = bits
            continue

        # ---- 新消息 ----
        if RE_HAS_MESSAGE_ID.search(name):
            current = MessageSpec(row_no=row_no)
            # type 先从参数名称找，例如「消息ID（type=1000）」「BMU消息ID(t=1050)」
            found = extract_type_ids(name)
            if not found:
                # 名称里没有就兜底看备注，例如「消息ID」+ 备注「type=1200(sc),1201(ka)」
                found = extract_type_ids(note)
            current.type_ids = found
            current.type_raw = " ".join(
                part for part in (name, note) if extract_type_ids(part or "")
            )
            slot.messages.append(current)
            current.fields.append(
                FieldSpec(name, bits, module, note, row_no, ROLE_HEADER)
            )
            continue

        # 消息ID之前不该有字段
        if current is None:
            continue

        # ---- 消息长度行：备注里藏着 length=NN字节 ----
        if name == "消息长度":
            if current.declared_bytes is None:
                length_match = RE_LENGTH.search(note)
                if length_match:
                    current.declared_bytes = int(length_match.group(1))
            current.fields.append(
                FieldSpec(name, bits, module, note, row_no, ROLE_HEADER)
            )
            continue

        # ---- 普通行 ----
        current.fields.append(
            FieldSpec(name, bits, module, note, row_no, guess_role(name, False))
        )

    return slot


def load_all_slots(xlsx_path: Optional[str] = None) -> List[SlotSpec]:
    """读整个工作簿，返回所有「通道N（slotM）」页；非通道页自动跳过。"""
    workbook = load_workbook(xlsx_path or default_spec_path(), read_only=True, data_only=True)

    slots: List[SlotSpec] = []
    for sheet_name in workbook.sheetnames:
        if not sheet_name.startswith(SHEET_CHANNEL_PREFIX):
            continue
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        slots.append(parse_sheet(rows, sheet_name))
    return slots
