# -*- coding: utf-8 -*-
"""
t_segment.py —— T 段解析：给 CSV 里的 T 段列配上中文名与判据说明。

T 段不需要位域解码：CSV 里的值已是解码好的中文枚举或数值（如「默认AB同时工作」、「2」），
只需按列序 1:1 贴到表里的行上，拿到中文名和判据说明。

    CSV 快遥  68 列 `TMXZJDT1xxx`  ↔  合并表「标准快遥」    68 行
    CSV 慢遥 111 列 `TMXZJDT2xxx`  ↔  合并表「标准慢遥51H」 111 行

映射靠"行序"而不是"名字匹配"（CSV 列名里只有序号、没有中文），
因此列数必须精确相等：对不上直接报错，不硬猜 —— 列错位一位则整张表列名全错。

T 段表第 3 列的列名叫「转换公式」，但内容是枚举说明与判据
（例如 `0000：正常; 1111：异常`），不可因列名而跳过。

用法：python t_segment.py --spec 打印两份 T 段定义；<csv> <行号> 打印该行的映射结果。
"""

from __future__ import annotations

import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from dataclasses import dataclass, field  # noqa: E402
from typing import Dict, List, Optional  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

from spec import default_spec_path  # noqa: E402

# --- T 段表的列号（从 1 开始）---
T_COL_NAME = 1   # 参数名称
T_COL_BITS = 2   # 参数位长
T_COL_NOTE = 3   # 列名叫「转换公式」，实际内容是枚举说明与判据

# --- CSV 里两种 T 段的列名前缀 ---
CSV_FAST_PREFIX = "TMXZJDT1"   # 快遥
CSV_SLOW_PREFIX = "TMXZJDT2"   # 慢遥

# --- 表里的 sheet 名 ---
SHEET_FAST = "标准快遥"
SHEET_SLOW_51H = "标准慢遥51H"

# 已知列数，用于校验（列数对不上就没有可靠的映射）
EXPECTED_COUNT = {SHEET_FAST: 68, SHEET_SLOW_51H: 111}


@dataclass
class TField:
    """T 段里的一个字段。"""

    index: int      # 第几列（从 0 开始）
    name: str       # 中文名
    bits: int       # 参数位长
    note: str       # 枚举说明 / 判据（表第 3 列）
    raw: str = ""   # CSV 里的原始文本（贴上值之后才有）


@dataclass
class TSegment:
    """一份 T 段定义（快遥 或 慢遥51H）。"""

    sheet_name: str
    fields: List[TField] = field(default_factory=list)

    @property
    def count(self) -> int:
        return len(self.fields)

    def with_values(self, values: List[str]) -> List[TField]:
        """把 CSV 一行的 T 段文本，按列序贴到字段上。

        Raises:
            ValueError: 列数与表行数不一致 —— 宁可不解析，也不错位。
        """
        if len(values) != self.count:
            raise ValueError(
                "列数不匹配：CSV 有 {} 列，{} 表里有 {} 行 —— 映射对不上，拒绝解析".format(
                    len(values), self.sheet_name, self.count
                )
            )
        return [
            TField(index=f.index, name=f.name, bits=f.bits, note=f.note, raw=values[f.index])
            for f in self.fields
        ]


def load_t_segments(spec_path: Optional[str] = None) -> Dict[str, TSegment]:
    """从合并表读出两份 T 段定义。

    Returns:
        {sheet 名: TSegment}，正常含「标准快遥」和「标准慢遥51H」
    """
    path = spec_path or default_spec_path()
    workbook = load_workbook(path, read_only=True, data_only=True)

    result: Dict[str, TSegment] = {}
    for sheet_name in (SHEET_FAST, SHEET_SLOW_51H):
        if sheet_name not in workbook.sheetnames:
            continue

        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        segment = TSegment(sheet_name=sheet_name)

        for row in rows[1:]:                       # 跳过表头
            name = str(row[T_COL_NAME - 1]).strip() if len(row) >= T_COL_NAME and row[T_COL_NAME - 1] is not None else ""
            if not name:
                continue                            # 空行跳过
            try:
                bits = int(row[T_COL_BITS - 1])
            except (TypeError, ValueError):
                bits = 0
            note = (
                str(row[T_COL_NOTE - 1]).strip()
                if len(row) >= T_COL_NOTE and row[T_COL_NOTE - 1] is not None
                else ""
            )
            segment.fields.append(
                TField(index=len(segment.fields), name=name, bits=bits, note=note)
            )

        result[sheet_name] = segment

    return result


def pick_segment(segments: Dict[str, TSegment], t_columns: List[str]) -> Optional[TSegment]:
    """看 CSV 的 T 段列名前缀，决定用哪份定义。

    快遥列名以 TMXZJDT1 开头，慢遥以 TMXZJDT2 开头。
    """
    if not t_columns:
        return None
    first = t_columns[0]
    if first.startswith(CSV_FAST_PREFIX):
        return segments.get(SHEET_FAST)
    if first.startswith(CSV_SLOW_PREFIX):
        return segments.get(SHEET_SLOW_51H)
    return None


def check_columns(segment: TSegment, columns: List[str]) -> str:
    """列数与表行数是否吻合；返回空串表示没问题。"""
    expected = EXPECTED_COUNT.get(segment.sheet_name)
    if expected is not None and len(columns) != expected:
        return "列数 {} ≠ {} 应有的 {} 列".format(len(columns), segment.sheet_name, expected)
    return ""


def one_line(text: str, width: int = 44) -> str:
    """把多行备注压成一行短摘要，便于打印。"""
    flat = " / ".join(part.strip() for part in (text or "").splitlines() if part.strip())
    return flat if len(flat) <= width else flat[: width - 1] + "…"


def print_spec(segments: Dict[str, TSegment]) -> None:
    """打印两份 T 段定义。"""
    for sheet_name, segment in segments.items():
        expected = EXPECTED_COUNT.get(sheet_name)
        flag = "OK" if expected == segment.count else "!! 应为 {} 行".format(expected)
        print("=" * 78)
        print("{}：{} 个字段   [{}]".format(sheet_name, segment.count, flag))
        print("=" * 78)
        for f in segment.fields:
            print("  第{:>3}列  {:>3}bit  {:<34} {}".format(
                f.index + 1, f.bits, f.name[:34], one_line(f.note)))
        print()


def dump_row(path: str, row_no: int, segments: Dict[str, TSegment]) -> None:
    """打印 CSV 某一行的 T 段映射结果。"""
    from table_csv import load_csv

    data = load_csv(path)
    if not data.t_columns:
        print("这个文件没有 T 段列")
        return
    if not 1 <= row_no <= len(data.rows):
        print("行号超出范围（1~{}）".format(len(data.rows)))
        return

    segment = pick_segment(segments, data.t_columns)
    if segment is None:
        print("认不出是哪种 T 段（首列 {}）".format(data.t_columns[0]))
        return

    problem = check_columns(segment, data.t_columns)
    row = data.rows[row_no - 1]

    print("=" * 78)
    print("文件 {}   第 {} 行   DMTime={}".format(path.replace("\\", "/").split("/")[-1], row_no, row.dm_time))
    print("T 段：{} 列，用表「{}」{} 行 → {}".format(
        len(data.t_columns), segment.sheet_name, segment.count,
        "列数吻合" if not problem else problem))
    print("=" * 78)

    if problem:
        return

    fields = segment.with_values(row.t_text)
    for f in fields:
        print("  {:>3}. {:<40} = {:<24} ｜ {}".format(
            f.index + 1, f.name[:40], str(f.raw)[:24], one_line(f.note, 56)))


def main() -> None:
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        return

    segments = load_t_segments()
    if not segments:
        print("没读到 T 段定义，检查合并表")
        return

    if args[0] == "--spec":
        print_spec(segments)
        return

    if len(args) < 2:
        print("用法: python t_segment.py <csv文件> <行号>")
        return

    dump_row(args[0], int(args[1]), segments)


if __name__ == "__main__":
    main()
